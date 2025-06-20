from docx import Document
from docx.shared import Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from datetime import datetime
import os
from docx.shared import RGBColor

RANKING_COLORS = {
    1: RGBColor(255, 215, 0),    # Gold (MVP)
    2: RGBColor(192, 192, 192),  # Silver (Desired)
    3: RGBColor(205, 127, 50),   # Bronze (Nice to Have)
}

def add_colored_paragraph(doc, text, ranking, style='Normal'):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    color = RANKING_COLORS.get(ranking, RGBColor(0, 0, 0))
    run.font.color.rgb = color
    return p



def add_table_of_contents(doc):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()

    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')

    instrText = OxmlElement('w:instrText')
    instrText.text = 'TOC \\o "1-3" \\h \\z \\u'

    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')

    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')

    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

    doc.add_paragraph().add_run("")


def int_to_letters(n):
    result = ''
    while n > 0:
        n -= 1
        result = chr(97 + (n % 26)) + result
        n //= 26
    return result


def export_feature_to_docx(feature, numbering=None, doc=None, include_notes=False, personnel_set=None):
    if numbering is None:
        numbering = []
    if doc is None:
        doc = Document()
        doc.add_heading("Feature Hierarchy Export", 0)
        add_table_of_contents(doc)

    if personnel_set is None:
        personnel_set = set()

    level = len(numbering)
    heading_number = '.'.join(map(str, numbering))

    # Feature heading
    heading_text = f"{heading_number} {feature.title} ({feature.ranking})"

    p = doc.add_heading(level=min(level + 1, 4))
    run = p.add_run(heading_text)
    color = RANKING_COLORS.get(feature.ranking, RGBColor(0, 0, 0))
    run.font.color.rgb = color
        
    # Feature description
    if feature.description:
        p = doc.add_paragraph(feature.description)
        p.style.font.size = Pt(10)

    # Add all feature images
    for attach in feature.attachments:
        if attach.filename.lower().endswith(('png', 'jpg', 'jpeg', 'gif')):
            image_path = os.path.join("static/uploads", attach.filename)
            if os.path.exists(image_path):
                try:
                    doc.add_picture(image_path, width=Inches(4.5))
                except Exception:
                    doc.add_paragraph(f"[Could not add feature image: {attach.filename}]")


    # Now user stories
    for i, story in enumerate(feature.user_stories, start=1):
        letter = int_to_letters(i)
        story_label = f"{heading_number}.{letter}"

        p = doc.add_paragraph(style='List Bullet')
        run = p.add_run(f"{story_label} {story.title} ({story.ranking})")
        color = RANKING_COLORS.get(story.ranking, RGBColor(0, 0, 0))
        run.font.color.rgb = color

        if story.created_by:
            personnel_set.add(story.created_by)
        if story.created_by or story.created_at:
            meta = f"  By: {story.created_by.name if story.created_by else 'Unknown'} | On: {story.created_at.strftime('%Y-%m-%d') if story.created_at else 'Unknown'}"
            doc.add_paragraph(meta, style='Caption')

        if story.assigned_to:
            assigned_text = f"Assigned to: {story.assigned_to.name}"
            p = doc.add_paragraph(assigned_text)
            p.style.font.size = Pt(10)

        if story.details:
            doc.add_paragraph(story.details, style='Intense Quote')

        if include_notes and story.notes:
            doc.add_paragraph("Notes:", style='List Continue')
            table = doc.add_table(rows=1, cols=3)
            hdr = table.rows[0].cells
            hdr[0].text = "Date"
            hdr[1].text = "Personnel"
            hdr[2].text = "Details"

            for note in story.notes:
                row = table.add_row().cells
                row[0].text = note.created_at.strftime('%Y-%m-%d')
                if note.personnel:
                    personnel_set.add(note.personnel)
                row[1].text = note.personnel.name if note.personnel else "Unknown"
                row[2].text = note.details

        for attachment in story.attachments:
            if attachment.filename.lower().endswith(('png', 'jpg', 'jpeg', 'gif')):
                image_path = os.path.join("static/uploads", attachment.filename)
                if os.path.exists(image_path):
                    try:
                        doc.add_picture(image_path, width=Inches(4.5))
                    except Exception:
                        doc.add_paragraph(f"[Could not add image: {attachment.filename}]")


    # recurse into subfeatures
    for i, subfeature in enumerate(feature.subfeatures, start=1):
        export_feature_to_docx(subfeature, numbering + [i], doc, include_notes, personnel_set)

    # Only add personnel list at root
    if numbering == []:
        doc.add_page_break()
        doc.add_heading("Personnel Involved", level=1)
        table = doc.add_table(rows=1, cols=3)
        hdr = table.rows[0].cells
        hdr[0].text = "Name"
        hdr[1].text = "Role"
        hdr[2].text = "Email"

        for person in sorted(personnel_set, key=lambda p: p.name.lower()):
            row = table.add_row().cells
            row[0].text = person.name
            row[1].text = person.role or ""
            email_run = row[2].paragraphs[0].add_run(person.email)
            email_run.hyperlink = f"mailto:{person.email}"

    return doc
